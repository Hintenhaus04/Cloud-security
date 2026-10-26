"""Maakt de (fictieve) SaxIT-dataset voor opdracht 4.1.

De dataset bestaat uit een data-inventaris van wat SaxIT opslaat, plus voorbeeldrecords van
klanten, CMDB-items en API-sleutels. Alle namen, adressen, IBAN's en sleutels zijn verzonnen.

Gebruik:
    pip install openpyxl
    python maak_dataset.py
"""

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path(__file__).with_name("saxit_dataset.xlsx")
random.seed(42)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# ID, data-element, beschrijving, voorbeeld, eigenaar, bron/systeem, persoonsgegevens, soort persoonsgegevens,
# volume, bewaartermijn
INVENTARIS = [
    ("D01", "Productbrochure CMDB", "Marketingmateriaal over de SaxIT CMDB", "brochure_saxit_cmdb_2026.pdf", "Marketing", "Website / S3", "Nee", "", "15 bestanden", "Tot vervanging"),
    ("D02", "Prijslijst abonnementen", "Standaardprijzen Basic, Pro, Enterprise", "prijslijst_2026.pdf", "Sales", "Website", "Nee", "", "1 bestand", "Tot vervanging"),
    ("D03", "Publieke API-documentatie", "Beschrijving van de REST API voor klanten", "api-docs/v1/openapi.yaml", "Development", "Developer portal", "Nee", "", "1 portal", "Per API-versie"),
    ("D04", "Persberichten en blogs", "Nieuws over releases en partners", "persbericht_lancering.docx", "Marketing", "Website", "Nee", "", "40 documenten", "Onbeperkt"),
    ("D05", "Algemene voorwaarden en privacyverklaring", "Juridische teksten op de website", "privacyverklaring_v3.pdf", "Legal", "Website", "Nee", "", "4 documenten", "Tot vervanging"),
    ("D06", "Statuspagina beschikbaarheid", "Uptime en geplande onderhoudsmomenten", "status.saxit.example", "Operations", "Statuspagina", "Nee", "", "Doorlopend", "90 dagen"),
    ("D07", "Werkinstructies en runbooks", "Hoe beheerders storingen oplossen en deployen", "runbook_database_failover.md", "Operations", "Wiki", "Nee", "", "120 pagina's", "Tot vervanging"),
    ("D08", "Architectuurplaten (high level)", "Overzicht van de 3-tier opzet zonder IP's of sleutels", "architectuur_overzicht.png", "Development", "Wiki", "Nee", "", "10 platen", "Tot vervanging"),
    ("D09", "Organigram en zakelijke contactgegevens medewerkers", "Naam, functie, zakelijk mailadres en telefoonnummer", "organigram_2026.xlsx", "HR", "Intranet", "Ja", "Gewone persoonsgegevens (zakelijk)", "85 medewerkers", "Duur dienstverband"),
    ("D10", "Projectplanningen en notulen", "Sprintplanningen, roadmaps, notulen van overleg", "roadmap_q1_2027.xlsx", "Management", "SharePoint", "Beperkt", "Namen deelnemers", "300 documenten", "7 jaar"),
    ("D11", "Infrastructure-as-code (zonder secrets)", "Terraform-scripts en pipeline-definities", "infra/main.tf", "Development", "GitHub", "Nee", "", "1 repository", "Onbeperkt"),
    ("D12", "Geanonimiseerde gebruiksstatistieken", "Aantal actieve gebruikers en API-calls per dag, niet herleidbaar", "usage_stats_2026-10.csv", "Product", "Datawarehouse", "Nee", "", "Dagelijks", "2 jaar"),
    ("D13", "Klantgegevens (contactpersonen)", "NAW, e-mail en telefoon van contactpersonen bij klanten", "zie tabblad Klanten", "Sales", "CRM / RDS", "Ja", "Gewone persoonsgegevens", "1.200 klanten", "Duur contract + 2 jaar"),
    ("D14", "Contracten en prijsafspraken", "Getekende contracten met kortingen en SLA's per klant", "contract_klant_0042.pdf", "Legal", "Documentbeheer", "Ja", "Naam en handtekening", "1.200 contracten", "7 jaar (fiscaal)"),
    ("D15", "Facturen en betaalgegevens", "Facturen met IBAN en betaalstatus", "factuur_2026_10_0042.pdf", "Finance", "Boekhouding", "Ja", "IBAN, naam", "15.000 per jaar", "7 jaar (fiscaal)"),
    ("D16", "CMDB-inhoud van klanten", "Configuratie-items van klanten: servers, IP's, softwareversies", "zie tabblad CMDB_items", "Klant (SaxIT is verwerker)", "CMDB / RDS", "Beperkt", "Namen beheerders", "2,4 mln items", "Duur contract"),
    ("D17", "API-sleutels en service-accounts", "Sleutels waarmee klanten en systemen de API aanroepen", "zie tabblad API_sleutels", "Security", "Secrets Manager", "Nee", "", "3.500 sleutels", "Tot intrekking"),
    ("D18", "Wachtwoordhashes eindgebruikers", "Gehashte wachtwoorden van gebruikers van de CMDB", "bcrypt-hash in tabel users", "Security", "RDS", "Ja", "Inloggegevens", "18.000 gebruikers", "Duur account"),
    ("D19", "Personeelsdossiers", "Arbeidscontracten, salaris, BSN, verzuim", "dossier_medewerker_017.pdf", "HR", "HR-systeem", "Ja", "BSN, salaris, gezondheid (bijzonder)", "85 dossiers", "2 jaar na uitdienst"),
    ("D20", "Pentestrapporten en kwetsbaarheden", "Gevonden kwetsbaarheden met exploitatiestappen", "pentest_rapport_2026.pdf", "Security", "Documentbeheer", "Nee", "", "6 rapporten", "5 jaar"),
    ("D21", "Audit- en toegangslogs", "Wie wanneer inlogde en welke data bekeek, met IP-adres", "cloudtrail/2026/10/26/*.json.gz", "Security", "CloudTrail / S3", "Ja", "IP-adres, gebruikers-ID", "5 GB per maand", "1 jaar"),
    ("D22", "Databaseback-ups", "Volledige snapshots van de productiedatabase", "rds-snapshot-2026-10-26", "Operations", "RDS / S3", "Ja", "Alles uit D13, D16, D18", "Dagelijks", "35 dagen"),
    ("D23", "Encryptiesleutels (keymateriaal)", "Eigen keymateriaal voor KMS (BYOK)", "PlaintextKeyMaterial.bin", "Security", "Offline kluis", "Nee", "", "3 sleutels", "Tot rotatie + archief"),
    ("D24", "Supporttickets", "Vragen en storingsmeldingen van klanten, soms met screenshots", "ticket_88231.eml", "Support", "Ticketsysteem", "Ja", "Naam, e-mail, inhoud melding", "9.000 per jaar", "2 jaar"),
]

VOORNAMEN = ["Sanne", "Daan", "Fleur", "Bram", "Noor", "Thijs", "Lotte", "Jesse", "Eva", "Ruben", "Iris", "Milan", "Femke", "Sem", "Julia"]
ACHTERNAMEN = ["de Vries", "Jansen", "Bakker", "Visser", "Smit", "Meijer", "de Boer", "Mulder", "Bos", "Vos", "Peters", "Hendriks", "Dekker", "Brouwer", "Kok"]
BEDRIJVEN = ["Veldhuis Logistiek", "Kortman Zorg", "Bureau Lindeboom", "Twentse Installatie", "Stadsdrukkerij Aalten", "Groenhof Retail", "Ijsseldelta Advies", "Maas & Co", "De Bruin ICT", "Huisartsen Oost", "Van Dam Bouw", "Noorderlicht Onderwijs", "Rijnland Transport", "Zuidkust Hotels", "Hofstede Accountants"]
PLAATSEN = ["Enschede", "Deventer", "Zwolle", "Hengelo", "Apeldoorn", "Arnhem", "Almelo", "Utrecht", "Groningen", "Nijmegen"]
STRATEN = ["Stationsstraat", "Kerkweg", "Hoofdstraat", "Molenlaan", "Industrieweg", "Beatrixlaan"]
ABONNEMENTEN = [("Basic", 49), ("Pro", 149), ("Enterprise", 499)]


def klanten():
    rows = []
    for i in range(15):
        voornaam, achternaam, bedrijf = VOORNAMEN[i], ACHTERNAMEN[i], BEDRIJVEN[i]
        domein = bedrijf.lower().replace(" & ", "").replace(" ", "") + ".example.nl"
        abonnement, prijs = random.choice(ABONNEMENTEN)
        korting = random.choice([0, 0, 5, 10, 15])
        rows.append((
            f"K{1001 + i}", bedrijf, f"{voornaam} {achternaam}",
            f"{voornaam.lower()}.{achternaam.lower().replace(' ', '')}@{domein}",
            f"06-{random.randint(10000000, 99999999)}",
            f"{random.choice(STRATEN)} {random.randint(1, 180)}", f"{random.randint(1000, 9999)} {random.choice('ABCDEFGHJKLMNPRSTVWXZ')}{random.choice('ABCDEFGHJKLMNPRSTVWXZ')}", random.choice(PLAATSEN),
            f"NL00TEST{random.randint(1000000000, 9999999999)}",
            abonnement, prijs, korting, f"2025-{random.randint(1, 12):02d}-{random.randint(1, 28):02d}",
        ))
    return rows


def cmdb_items():
    rows = []
    typen = [("Server", ["Ubuntu 22.04", "Windows Server 2019", "Windows Server 2022", "RHEL 8"]),
             ("Firewall", ["FortiOS 7.2.4", "pfSense 2.7"]),
             ("Database", ["MySQL 8.0.32", "SQL Server 2019", "PostgreSQL 14"]),
             ("Werkplek", ["Windows 11 23H2", "Windows 10 22H2"])]
    for i in range(15):
        klant = f"K{1001 + random.randint(0, 14)}"
        soort, versies = random.choice(typen)
        rows.append((
            f"CI-{20000 + i}", klant, f"{soort.lower()[:3]}-{random.choice(['prod', 'acc', 'kant'])}-{random.randint(1, 30):02d}",
            soort, random.choice(versies), f"10.{random.randint(0, 50)}.{random.randint(0, 255)}.{random.randint(2, 254)}",
            random.choice(["Rack A3", "Azure West-EU", "Kantoor 1e verdieping", "Serverruimte"]),
            f"{random.choice(VOORNAMEN)} {random.choice(ACHTERNAMEN)}",
            random.choice(["Actief", "Actief", "Actief", "Uitgefaseerd"]),
            random.choice(["", "", "Patch achterstand", "End-of-life", ""]),
        ))
    return rows


def api_sleutels():
    rows = []
    for i in range(10):
        key = "sxt_live_" + "".join(random.choice("abcdefghijklmnopqrstuvwxyz0123456789") for _ in range(24))
        rows.append((
            f"AK-{300 + i}", f"K{1001 + random.randint(0, 14)}", key[:13] + "…" + key[-4:],
            random.choice(["read", "read", "read,write", "admin"]),
            f"2026-{random.randint(1, 10):02d}-{random.randint(1, 28):02d}",
            random.choice(["Actief", "Actief", "Ingetrokken"]),
        ))
    return rows


def write_sheet(wb, title, header, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for row in rows:
        ws.append(row)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    wb = Workbook()
    info = wb.active
    info.title = "Toelichting"
    for line in [
        ["SaxIT dataset voor opdracht 4.1 (dataclassificatie)"],
        [],
        ["Deze dataset is fictief en voor deze opdracht gemaakt. Namen, adressen, IBAN's en sleutels bestaan niet."],
        ["Tabblad Data-inventaris: alle soorten data die SaxIT opslaat en verwerkt (24 data-elementen)."],
        ["Tabbladen Klanten, CMDB_items en API_sleutels: voorbeeldrecords uit de productiedatabase."],
    ]:
        info.append(line)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 110

    write_sheet(wb, "Data-inventaris",
                ["ID", "Data-element", "Beschrijving", "Voorbeeld", "Eigenaar", "Bron / systeem",
                 "Persoonsgegevens", "Soort persoonsgegevens", "Volume", "Bewaartermijn"],
                INVENTARIS, [6, 34, 48, 30, 18, 18, 14, 30, 16, 20])
    write_sheet(wb, "Klanten",
                ["Klant-ID", "Bedrijf", "Contactpersoon", "E-mail", "Telefoon", "Adres", "Postcode", "Plaats",
                 "IBAN", "Abonnement", "Prijs p/m (EUR)", "Korting %", "Klant sinds"],
                klanten(), [10, 22, 18, 40, 13, 20, 10, 12, 22, 12, 14, 10, 12])
    write_sheet(wb, "CMDB_items",
                ["CI-ID", "Klant-ID", "Hostnaam", "Type", "Versie", "IP-adres", "Locatie", "Beheerder", "Status", "Opmerking"],
                cmdb_items(), [10, 10, 16, 12, 20, 15, 22, 18, 13, 18])
    write_sheet(wb, "API_sleutels",
                ["Key-ID", "Klant-ID", "Sleutel (ingekort)", "Rechten", "Aangemaakt", "Status"],
                api_sleutels(), [10, 10, 26, 12, 12, 12])

    wb.save(OUT)
    print(f"Opgeslagen: {OUT}")


if __name__ == "__main__":
    main()
